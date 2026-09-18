#!/usr/bin/env python3
"""Ejecuta 3 escenarios ficticios GEMEPH sobre datos oficiales EPH y actualiza entregables."""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from gemeph.export import export_word_bytes
from gemeph.panel import load_or_build_panel, periodo_texto
from gemeph.scenario import compare_rows, lever_baselines, run_scenario

OUT = Path("/Users/claudiolarrea/Library/CloudStorage/OneDrive-Personal/11 Investigacion/2026/GEMEPH")
RESULTS = OUT / "02_Resultados_empiricos"
ANEXO = OUT / "01_Proyecto_investigacion" / "Anexo_III_Informe_Final_GEMEPH.docx"
ANEXO_ROOT = OUT / "Anexo_III_Informe_Final_GEMEPH.docx"
ANEXO_DL = Path("/Users/claudiolarrea/Downloads/Anexo_III_Informe_Final_GEMEPH.docx")

# Tres escenarios ficticios contrastantes (what-if), anclados al baseline oficial
SCENARIOS: list[dict[str, Any]] = [
    {
        "id": "E1",
        "nombre": "Conectividad inclusiva (cierre de brecha digital en quintil bajo)",
        "narrativa": (
            "Escenario ficticio de política de conectividad: se supone que casi todos los hogares "
            "del quintil de ingresos más bajo pasan a tener internet en el hogar."
        ),
        "targets_fn": lambda b: {
            "internet_quintil_i": 99.0,
            "pct_superior": b["pct_superior"],
            "pct_empleo_formal": b["pct_empleo_formal"],
        },
    },
    {
        "id": "E2",
        "nombre": "Expansión educativa superior",
        "narrativa": (
            "Escenario ficticio educativo: se eleva la proporción de personas con educación "
            "universitaria completa hasta un umbral de alta cobertura, manteniendo conectividad y formalidad en valores observados."
        ),
        "targets_fn": lambda b: {
            "internet_quintil_i": b.get("internet_quintil_i", b["pct_superior"]),
            "pct_superior": 40.0,
            "pct_empleo_formal": b["pct_empleo_formal"],
        },
    },
    {
        "id": "E3",
        "nombre": "Formalización laboral + educación (paquete integrado)",
        "narrativa": (
            "Escenario ficticio integrado: combina mayor formalización del empleo entre ocupados "
            "con un salto moderado-alto en educación superior, sin alterar la meta de internet del quintil bajo más allá de una mejora parcial."
        ),
        "targets_fn": lambda b: {
            "internet_quintil_i": min(99.0, float(b.get("internet_quintil_i", 90)) + 3.0),
            "pct_superior": 30.0,
            "pct_empleo_formal": 90.0,
        },
    },
]


def _fmt(v: Any, nd: int = 4) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "—"
    if isinstance(v, float):
        return f"{v:.{nd}f}"
    return str(v)


def _kpi_row(label: str, source: str, kpis: dict[str, Any], modelo: dict[str, Any] | None = None) -> dict[str, Any]:
    return {
        "escenario": label,
        "fuente": source,
        "n_individuos": kpis.get("n_individuos"),
        "idx_exclusion_digital": kpis.get("idx_exclusion_digital"),
        "pct_exclusion_digital_alta": kpis.get("pct_exclusion_digital_alta"),
        "vulnerabilidad_social": kpis.get("vulnerabilidad_social"),
        "score_movilidad_proxy": kpis.get("score_movilidad_proxy"),
        "pct_superior": kpis.get("pct_superior"),
        "pct_ocupado": kpis.get("pct_ocupado"),
        "pct_informal_ocupados": kpis.get("pct_informal_ocupados"),
        "pct_exclusion_predicha_modelo": (modelo or {}).get("pct_exclusion_predicho_base")
        if source == "oficial"
        else (modelo or {}).get("pct_exclusion_predicho_escenario"),
    }


def build_excel(
    baseline_kpis: dict[str, Any],
    levers: dict[str, float],
    runs: list[dict[str, Any]],
    meta: dict[str, Any],
    cat_path: Path | None = None,
) -> bytes:
    import io

    buf = io.BytesIO()
    # Comparación consolidada
    rows = [_kpi_row("Baseline oficial EPH", "oficial", baseline_kpis, None)]
    # attach model base from first run
    if runs:
        rows[0]["pct_exclusion_predicha_modelo"] = runs[0]["result"]["modelo"].get("pct_exclusion_predicho_base")

    for r in runs:
        rows.append(
            _kpi_row(
                f"{r['id']} — {r['nombre']}",
                "escenario_ficticio",
                r["result"]["scenario_kpis"],
                r["result"]["modelo"],
            )
        )
    cmp_df = pd.DataFrame(rows)

    delta_rows = []
    for r in runs:
        d = r["result"]["deltas"]
        delta_rows.append(
            {
                "escenario_id": r["id"],
                "escenario": r["nombre"],
                **{f"delta_{k}": v for k, v in d.items()},
                "pred_exclusion_base_%": r["result"]["modelo"].get("pct_exclusion_predicho_base"),
                "pred_exclusion_escenario_%": r["result"]["modelo"].get("pct_exclusion_predicho_escenario"),
                "delta_pred_exclusion_pp": (
                    None
                    if r["result"]["modelo"].get("pct_exclusion_predicho_base") is None
                    or r["result"]["modelo"].get("pct_exclusion_predicho_escenario") is None
                    else round(
                        float(r["result"]["modelo"]["pct_exclusion_predicho_escenario"])
                        - float(r["result"]["modelo"]["pct_exclusion_predicho_base"]),
                        2,
                    )
                ),
            }
        )

    palancas = []
    for r in runs:
        palancas.append({"escenario_id": r["id"], "escenario": r["nombre"], "tipo": "baseline_oficial", **levers})
        palancas.append({"escenario_id": r["id"], "escenario": r["nombre"], "tipo": "meta_ficticia", **r["targets"]})

    narrativa = [
        {
            "escenario_id": r["id"],
            "nombre": r["nombre"],
            "narrativa": r["narrativa"],
            "interpretacion": r["interpretacion"],
        }
        for r in runs
    ]

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame([{"campo": k, "valor": str(v)} for k, v in meta.items()]).to_excel(
            writer, sheet_name="metadatos", index=False
        )
        pd.DataFrame([{"palanca": k, "valor_oficial_%": v} for k, v in levers.items()]).to_excel(
            writer, sheet_name="baseline_palancas_oficiales", index=False
        )
        cmp_df.to_excel(writer, sheet_name="comparacion_baseline_vs_3E", index=False)
        pd.DataFrame(delta_rows).to_excel(writer, sheet_name="deltas_escenarios", index=False)
        pd.DataFrame(palancas).to_excel(writer, sheet_name="palancas_oficial_vs_meta", index=False)
        pd.DataFrame(narrativa).to_excel(writer, sheet_name="narrativa_escenarios", index=False)
        for r in runs:
            compare_rows(r["result"]).to_excel(writer, sheet_name=f"detalle_{r['id']}"[:31], index=False)
        # Reusar catálogo territorial si existe en excel previo
        prev = RESULTS / "GEMEPH_resultados_2022_2024_T4_tic_20260821.xlsx"
        if prev.exists():
            try:
                cat = pd.read_excel(prev, sheet_name="catalogo_31_aglomerados")
                cat.to_excel(writer, sheet_name="catalogo_31_aglomerados", index=False)
            except Exception:
                pass
            try:
                hall = pd.read_excel(prev, sheet_name="hallazgos_OE1_OE6")
                hall.to_excel(writer, sheet_name="hallazgos_OE1_OE6", index=False)
            except Exception:
                pass
    buf.seek(0)
    return buf.getvalue()


def build_word(
    meta: dict[str, Any],
    levers: dict[str, float],
    baseline_kpis: dict[str, Any],
    runs: list[dict[str, Any]],
) -> bytes:
    import io

    doc = Document()
    doc.add_heading(meta["titulo"], level=0)
    doc.add_paragraph(f"Período (datos oficiales): {meta['periodo']}")
    doc.add_paragraph(f"Fuente: {meta['fuente']}")
    doc.add_paragraph(
        "Propósito del gemelo digital: exponer los datos oficiales de la EPH a escenarios ficticios (what-if) "
        "para observar qué ocurriría en indicadores de exclusión digital, vulnerabilidad, movilidad proxy, "
        "educación y empleo si se alteraran palancas de política."
    )
    doc.add_paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    doc.add_heading("1. Baseline oficial (EPH-INDEC)", level=1)
    doc.add_paragraph(
        f"Registros: {baseline_kpis.get('n_individuos'):,} · "
        f"Exclusión digital: {_fmt(baseline_kpis.get('idx_exclusion_digital'))} · "
        f"% exclusión alta: {_fmt(baseline_kpis.get('pct_exclusion_digital_alta'), 2)} · "
        f"Vulnerabilidad: {_fmt(baseline_kpis.get('vulnerabilidad_social'))} · "
        f"Movilidad proxy: {_fmt(baseline_kpis.get('score_movilidad_proxy'))} · "
        f"Superior: {_fmt(baseline_kpis.get('pct_superior'), 2)}% · "
        f"Ocupación: {_fmt(baseline_kpis.get('pct_ocupado'), 2)}% · "
        f"Informalidad (ocupados): {_fmt(baseline_kpis.get('pct_informal_ocupados'), 2)}%."
    )
    doc.add_paragraph(
        "Palancas oficiales observadas: "
        + "; ".join(f"{k}={v}%" for k, v in levers.items())
        + "."
    )

    doc.add_heading("2. Tres escenarios ficticios", level=1)
    for r in runs:
        doc.add_heading(f"{r['id']}. {r['nombre']}", level=2)
        doc.add_paragraph(r["narrativa"])
        doc.add_paragraph(
            "Metas ficticias: " + "; ".join(f"{k}={v}%" for k, v in r["targets"].items()) + "."
        )
        doc.add_paragraph("Resultados contrafactuales (respecto del baseline oficial):")
        for _, row in compare_rows(r["result"]).iterrows():
            doc.add_paragraph(
                f"- {row['Indicador']}: oficial {_fmt(row['Baseline'])} → escenario {_fmt(row['Escenario'])} "
                f"(Δ {_fmt(row.get('Cambio'))})"
            )
        modelo = r["result"]["modelo"]
        doc.add_paragraph(
            f"Modelo predictivo de exclusión alta: {_fmt(modelo.get('pct_exclusion_predicho_base'), 2)}% "
            f"→ {_fmt(modelo.get('pct_exclusion_predicho_escenario'), 2)}%."
        )
        doc.add_paragraph(r["interpretacion"])

    doc.add_heading("3. Lectura conjunta", level=1)
    doc.add_paragraph(
        "Los tres escenarios ilustran el valor del gemelo: el baseline permanece anclado a microdatos oficiales, "
        "mientras que cada simulación permite anticipar magnitudes de cambio ante intervenciones hipotéticas. "
        "E1 aísla el canal digital en hogares de bajos ingresos; E2 el canal educativo; E3 un paquete integrado "
        "educación–formalización. Las diferencias de impacto entre escenarios informan prioridades de política "
        "sin confundirse con proyecciones oficiales del INDEC."
    )
    doc.add_paragraph(
        "Nota metodológica: los escenarios son contrafactuales construidos sobre el panel EPH; no constituyen "
        "pronósticos oficiales ni causalidad identificada."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _set_para(p, text: str, bold: bool = False) -> None:
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def _insert_after(paragraph, text: str, *, bold: bool = False, space_after: int = 6):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    para = Paragraph(new_p, paragraph._parent)
    para.paragraph_format.space_after = Pt(space_after)
    if text:
        _set_para(para, text, bold=bold)
    return para


def update_anexo_iii(runs: list[dict[str, Any]], baseline_kpis: dict[str, Any], levers: dict[str, float], meta: dict[str, Any]) -> None:
    paths = [ANEXO, ANEXO_ROOT, ANEXO_DL]
    bloque = _texto_escenarios(runs, baseline_kpis, levers, meta)

    for path in paths:
        if not path.exists():
            continue
        doc = Document(str(path))

        # Replace / enrich Resultados y discusión paragraph
        replaced = False
        for p in doc.paragraphs:
            t = p.text.strip()
            if t.startswith("Resultados y discusión") or (
                "Resultados nacionales (ponderados)" in t and "exclusión digital" in t.lower()
            ):
                _set_para(p, "Resultados y discusión:\n" + bloque)
                replaced = True
                break

        if not replaced:
            # Insert before Conclusiones y proyecciones
            for p in doc.paragraphs:
                if p.text.strip().startswith("Conclusiones y proyecciones"):
                    # insert before by adding previous content via replacing nearby
                    cursor = p
                    # Use addprevious chain in reverse
                    from docx.oxml import OxmlElement as OE

                    new_p = OE("w:p")
                    cursor._p.addprevious(new_p)
                    para = Paragraph(new_p, cursor._parent)
                    _set_para(para, "Resultados y discusión — escenarios ficticios del gemelo:\n" + bloque)
                    break

        # Enrich síntesis if too short / without escenarios
        for p in doc.paragraphs:
            if p.text.strip().startswith("Síntesis de los resultados"):
                if "escenarios ficticios" not in p.text.lower() and "E1" not in p.text:
                    extra = (
                        f" El gemelo se usó para exponer el baseline oficial a tres escenarios ficticios: "
                        f"E1 conectividad inclusiva, E2 expansión educativa superior y E3 formalización+educación. "
                        f"Baseline: exclusión={_fmt(baseline_kpis.get('idx_exclusion_digital'))}, "
                        f"vulnerabilidad={_fmt(baseline_kpis.get('vulnerabilidad_social'))}, "
                        f"movilidad proxy={_fmt(baseline_kpis.get('score_movilidad_proxy'))}."
                    )
                    _set_para(p, p.text.rstrip() + extra)
                break

        # Enrich conclusiones
        for p in doc.paragraphs:
            if p.text.strip().startswith("Principales conclusiones"):
                if "E1" not in p.text:
                    _set_para(
                        p,
                        p.text.rstrip()
                        + " 6) El valor central de GEMEPH se demuestra al contrastar datos oficiales con tres escenarios "
                        "ficticios (E1 digital, E2 educativo, E3 integrado), registrando deltas en exclusión, vulnerabilidad, "
                        "movilidad proxy y predicción de exclusión alta.",
                    )
                break

        doc.save(str(path))
        print("Anexo actualizado:", path)


def _texto_escenarios(runs, baseline_kpis, levers, meta) -> str:
    parts = [
        f"Datos oficiales (EPH-INDEC, {meta['periodo']}, n={baseline_kpis.get('n_individuos'):,}): "
        f"exclusión digital={_fmt(baseline_kpis.get('idx_exclusion_digital'))}; "
        f"% exclusión alta={_fmt(baseline_kpis.get('pct_exclusion_digital_alta'), 2)}; "
        f"vulnerabilidad={_fmt(baseline_kpis.get('vulnerabilidad_social'))}; "
        f"movilidad proxy={_fmt(baseline_kpis.get('score_movilidad_proxy'))}; "
        f"superior={_fmt(baseline_kpis.get('pct_superior'), 2)}%; "
        f"ocupación={_fmt(baseline_kpis.get('pct_ocupado'), 2)}%; "
        f"informalidad ocupados={_fmt(baseline_kpis.get('pct_informal_ocupados'), 2)}%. "
        f"Palancas oficiales: " + ", ".join(f"{k}={v}%" for k, v in levers.items()) + ".",
        "El gemelo digital no reemplaza al dato oficial: lo toma como baseline y lo expone a escenarios ficticios (what-if) para estimar qué ocurriría ante intervenciones hipotéticas.",
    ]
    for r in runs:
        d = r["result"]["deltas"]
        modelo = r["result"]["modelo"]
        parts.append(
            f"{r['id']} — {r['nombre']}: {r['narrativa']} "
            f"Metas: " + ", ".join(f"{k}={v}%" for k, v in r["targets"].items()) + ". "
            f"Resultados: Δ exclusión={_fmt(d.get('idx_exclusion_digital'))}; "
            f"Δ % exclusión alta={_fmt(d.get('pct_exclusion_digital_alta'), 2)}; "
            f"Δ vulnerabilidad={_fmt(d.get('vulnerabilidad_social'))}; "
            f"Δ movilidad proxy={_fmt(d.get('score_movilidad_proxy'))}; "
            f"Δ superior={_fmt(d.get('pct_superior'), 2)} pp. "
            f"Predicción exclusión alta: {_fmt(modelo.get('pct_exclusion_predicho_base'), 2)}% → "
            f"{_fmt(modelo.get('pct_exclusion_predicho_escenario'), 2)}%. "
            f"{r['interpretacion']}"
        )
    parts.append(
        "Discusión: E1, E2 y E3 muestran sensibilidades distintas del sistema sociodemográfico. "
        "La comparación entre canales (digital, educativo e integrado) permite priorizar hipótesis de política "
        "sin confundir simulación con proyección oficial del INDEC."
    )
    return " ".join(parts)


def interpretar(r_id: str, result: dict[str, Any]) -> str:
    d = result["deltas"]
    if r_id == "E1":
        return (
            f"Interpretación: al forzar conectividad casi universal en el quintil bajo, el cambio en el índice de exclusión "
            f"digital es {_fmt(d.get('idx_exclusion_digital'))} y en % de exclusión alta {_fmt(d.get('pct_exclusion_digital_alta'), 2)} pp. "
            "El escenario aísla el canal digital sobre la base oficial."
        )
    if r_id == "E2":
        return (
            f"Interpretación: el salto educativo (superior) produce Δ superior={_fmt(d.get('pct_superior'), 2)} pp y "
            f"Δ movilidad proxy={_fmt(d.get('score_movilidad_proxy'))}, evidenciando el canal educativo del gemelo."
        )
    return (
        f"Interpretación: el paquete integrado mueve simultáneamente educación (Δ superior={_fmt(d.get('pct_superior'), 2)} pp) "
        f"y formalidad/condiciones laborales asociadas, con Δ vulnerabilidad={_fmt(d.get('vulnerabilidad_social'))} "
        f"y Δ movilidad proxy={_fmt(d.get('score_movilidad_proxy'))}."
    )


def main() -> None:
    years = [2022, 2023, 2024]
    trimestre = 4
    modulo = "tic"
    print("==> Cargando panel oficial…")
    panel, val, run_id = load_or_build_panel(years, trimestre, modulo=modulo, force_download=False)
    periodo = periodo_texto(years, trimestre)
    print(f"    n={len(panel):,} run_id={run_id}")

    levers = lever_baselines(panel, include_tic=True)
    print("    levers oficiales:", levers)

    # Baseline KPIs via first scenario's baseline
    runs: list[dict[str, Any]] = []
    for spec in SCENARIOS:
        targets = spec["targets_fn"](levers)
        # Ensure only known levers
        targets = {k: float(v) for k, v in targets.items() if k in levers or k in ("internet_quintil_i", "pct_superior", "pct_empleo_formal")}
        print(f"==> {spec['id']}: {targets}")
        result = run_scenario(panel, targets, include_tic=True)
        runs.append(
            {
                "id": spec["id"],
                "nombre": spec["nombre"],
                "narrativa": spec["narrativa"],
                "targets": targets,
                "result": result,
                "interpretacion": "",  # fill below
            }
        )
        runs[-1]["interpretacion"] = interpretar(spec["id"], result)

    baseline_kpis = runs[0]["result"]["baseline_kpis"]
    meta = {
        "titulo": "GEMEPH — Datos oficiales EPH vs 3 escenarios ficticios",
        "periodo": periodo,
        "modulo": modulo,
        "run_id": run_id,
        "fuente": "INDEC — EPH (hogar + individuo + TIC)",
        "registros_panel": len(panel),
        "n_escenarios": 3,
        "plataforma": "https://eph-analyzer.streamlit.app/GEMEPH",
        "generado": datetime.now().isoformat(timespec="seconds"),
        "nota": "Los escenarios son contrafactuales; el baseline es dato oficial.",
    }

    stamp = datetime.now().strftime("%Y%m%d")
    RESULTS.mkdir(parents=True, exist_ok=True)
    excel_path = RESULTS / f"GEMEPH_resultados_3escenarios_{run_id}_{stamp}.xlsx"
    word_path = RESULTS / f"GEMEPH_informe_3escenarios_{run_id}_{stamp}.docx"
    json_path = RESULTS / f"GEMEPH_resumen_3escenarios_{run_id}_{stamp}.json"

    print("==> Excel…")
    excel_path.write_bytes(build_excel(baseline_kpis, levers, runs, meta))
    print("==> Word…")
    word_path.write_bytes(build_word(meta, levers, baseline_kpis, runs))

    # Copias en raíz GEMEPH y Downloads
    for dest_dir in [OUT, Path("/Users/claudiolarrea/Downloads")]:
        (dest_dir / excel_path.name).write_bytes(excel_path.read_bytes())
        (dest_dir / word_path.name).write_bytes(word_path.read_bytes())

    payload = {
        "meta": meta,
        "levers_oficiales": levers,
        "baseline_kpis": baseline_kpis,
        "escenarios": [
            {
                "id": r["id"],
                "nombre": r["nombre"],
                "narrativa": r["narrativa"],
                "targets": r["targets"],
                "deltas": r["result"]["deltas"],
                "scenario_kpis": r["result"]["scenario_kpis"],
                "modelo": r["result"]["modelo"],
                "interpretacion": r["interpretacion"],
            }
            for r in runs
        ],
        "archivos": {"excel": str(excel_path), "word": str(word_path)},
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    print("==> Anexo III…")
    update_anexo_iii(runs, baseline_kpis, levers, meta)

    print("LISTO")
    print(excel_path)
    print(word_path)
    for r in runs:
        print(r["id"], r["result"]["deltas"])


if __name__ == "__main__":
    main()
